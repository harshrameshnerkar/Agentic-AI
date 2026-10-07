"""
pdf_loader.py
=============
Document Ingestion & Metadata Extraction Module (Day 3 - Session 2)

Learning Goals Covered:
- Ingesting PDFs and extracting clean text per page
- Preserving critical metadata (source filename, page number, total pages, character count)
- Architecture patterns for loading DOCX and Web pages into normalized Document schemas.

Metadata Preservation Best Practice:
Never concatenate entire PDFs into one big raw string. Splitting per page and preserving
source metadata enables downstream citations (e.g., "See Cloud SLA, Page 3") and
prevents cross-page context contamination.
"""

from pathlib import Path
from typing import List, Dict, Any
from pypdf import PdfReader


class Document:
    """
    Standard normalized document structure compatible with modern RAG frameworks.
    Encapsulates text content and page-level metadata.
    """
    def __init__(self, page_content: str, metadata: Dict[str, Any]):
        self.page_content = page_content
        self.metadata = metadata

    def __repr__(self) -> str:
        source = self.metadata.get("source", "unknown")
        page = self.metadata.get("page", "?")
        chars = len(self.page_content)
        return f"<Document source='{source}' page={page} chars={chars}>"


def load_pdf_with_metadata(pdf_path: Path) -> List[Document]:
    """
    Loads a PDF file using pypdf, extracting text on a per-page basis
    while attaching structured metadata.

    Args:
        pdf_path: Path to the PDF file.

    Returns:
        List of Document objects, one for each non-empty page.
    """
    if not pdf_path.exists():
        raise FileNotFoundError(f"PDF file not found: {pdf_path}")

    reader = PdfReader(str(pdf_path))
    total_pages = len(reader.pages)
    documents: List[Document] = []

    for page_idx, page in enumerate(reader.pages):
        page_num = page_idx + 1
        text = page.extract_text() or ""
        text = text.strip()

        if not text:
            continue

        metadata = {
            "source": pdf_path.name,
            "file_path": str(pdf_path.resolve()),
            "page": page_num,
            "total_pages": total_pages,
            "char_count": len(text),
        }
        documents.append(Document(page_content=text, metadata=metadata))

    return documents


def load_all_pdfs(docs_directory: Path) -> List[Document]:
    """
    Discovers and loads all PDF files in the specified directory.

    Args:
        docs_directory: Path to folder containing PDFs.

    Returns:
        List of all extracted Document objects across all PDFs.
    """
    pdf_files = sorted(list(docs_directory.glob("*.pdf")))
    if not pdf_files:
        raise ValueError(f"No PDF files found in {docs_directory}")

    all_docs: List[Document] = []
    for pdf_file in pdf_files:
        docs = load_pdf_with_metadata(pdf_file)
        all_docs.extend(docs)
        print(f"  [PDF Loaded] {pdf_file.name}: {len(docs)} pages extracted.")

    return all_docs


# ---------------------------------------------------------------------------
# EDUCATIONAL EXTENSION PATTERNS: DOCX & WEB PAGE INGESTION
# ---------------------------------------------------------------------------
def how_to_load_docx_example(docx_path: Path) -> str:
    """
    Conceptual guide for ingesting Microsoft Word (.docx) files.
    Requires: pip install python-docx
    
    Implementation pattern:
    ```python
    import docx

    doc = docx.Document(docx_path)
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    full_text = "\n\n".join(paragraphs)
    # Metadata includes author, created date, and file properties
    ```
    """
    return "Use python-docx to iterate over doc.paragraphs and doc.tables, capturing sections."


def how_to_load_webpage_example(url: str) -> str:
    """
    Conceptual guide for ingesting HTML/Web pages.
    Requires: pip install beautifulsoup4 requests (or langchain-community WebBaseLoader)

    Implementation pattern:
    ```python
    import requests
    from bs4 import BeautifulSoup

    resp = requests.get(url, timeout=10)
    soup = BeautifulSoup(resp.text, "html.parser")
    # Clean scripts, styles, navigation, footer
    for tag in soup(["script", "style", "nav", "footer", "header"]):
        tag.decompose()
    clean_text = soup.get_text(separator="\n").strip()
    metadata = {"source": url, "title": soup.title.string if soup.title else ""}
    ```
    """
    return "Use BeautifulSoup to strip boilerplate (nav/footer/scripts) and retain core article content."


if __name__ == "__main__":
    docs_dir = Path(__file__).parent / "docs"
    print("Testing PDF Ingestion on directory:", docs_dir)
    docs = load_all_pdfs(docs_dir)
    print(f"\nTotal pages ingested: {len(docs)}")
    for d in docs[:3]:
        print(f" - {d} | Preview: {d.page_content[:90]}...")
